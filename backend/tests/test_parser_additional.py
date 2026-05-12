import io
import pytest
import fitz
import docx
import openpyxl

from backend.services.parser import parse_and_chunk, CHUNK_SIZE, _sanitize, _chunk

def _make_pdf(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), text)
    return doc.tobytes()

def _make_docx(text: str) -> bytes:
    document = docx.Document()
    document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()

def _make_xlsx(rows: list[list[str]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

# Existing tests (preserved) ---------------------------------------------------

def test_parse_docx_strips_control_chars():
    raw = "valid\x00\x08content"
    docx_bytes = _make_docx(raw)
    chunks = parse_and_chunk(docx_bytes, "doc.docx")
    combined = "".join(chunks)
    assert "\x00" not in combined
    assert "\x08" not in combined
    assert "validcontent" in combined

def test_chunk_empty_input():
    result = _chunk("")
    assert result == [""]

def test_chunk_exact_limit():
    text = "c" * CHUNK_SIZE
    result = _chunk(text)
    assert len(result) == 1
    assert len(result[0]) == CHUNK_SIZE

def test_sanitize_full_control_range():
    control_chars = ''.join(chr(i) for i in list(range(0x00, 0x20)) + [0x7f])
    sanitized = _sanitize(control_chars)
    allowed = {"\t", "\n", "\r"}
    for ch in sanitized:
        assert ch in allowed

# New tests -------------------------------------------------------------------

def test_parse_pdf_extraction_and_sanitize():
    raw = "PDF\x00content\x07test"
    pdf_bytes = _make_pdf(raw)
    chunks = parse_and_chunk(pdf_bytes, "file.pdf")
    combined = "".join(chunks)
    assert "PDF" in combined
    assert "content" in combined
    assert "test" in combined
    assert "\x00" not in combined
    assert "\x07" not in combined

def test_parse_xlsx_extraction_and_sanitize():
    rows = [["A\x01", "B\x7f"], ["C", "D"]]
    xlsx_bytes = _make_xlsx(rows)
    chunks = parse_and_chunk(xlsx_bytes, "sheet.xlsx")
    combined = "".join(chunks)
    # values should be tab‑separated and newline‑separated
    assert "A" in combined and "B" in combined and "C" in combined and "D" in combined
    assert "\x01" not in combined and "\x7f" not in combined

def test_chunk_just_over_limit():
    text = "x" * (CHUNK_SIZE + 10)
    result = _chunk(text)
    assert len(result) == 2
    assert len(result[0]) == CHUNK_SIZE
    assert len(result[1]) == 10
