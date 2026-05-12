import io
import pytest
import fitz
import docx
import openpyxl

from backend.services.parser import parse_and_chunk, CHUNK_SIZE, _sanitize, _chunk


# ---------------------------------------------------------------------------
# Helpers to build in-memory fixture documents
# ---------------------------------------------------------------------------

def _make_pdf(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), text)
    return doc.tobytes()


def _make_docx(text: str) -> bytes:
    document = docx.Document()
    document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _make_xlsx(rows: list[list[str]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Format parsing: each format yields the correct text
# ---------------------------------------------------------------------------

def test_parse_txt_returns_text():
    data = "hello world".encode("utf-8")
    chunks = parse_and_chunk(data, "sample.txt")
    assert "".join(chunks) == "hello world"


def test_parse_txt_utf8_content():
    text = "café résumé"
    chunks = parse_and_chunk(text.encode("utf-8"), "note.TXT")
    assert "".join(chunks) == text


def test_parse_pdf_returns_text():
    content = "Invoice total: 500"
    pdf_bytes = _make_pdf(content)
    chunks = parse_and_chunk(pdf_bytes, "doc.pdf")
    combined = "".join(chunks)
    assert content in combined


def test_parse_docx_returns_text():
    content = "Contract clause one"
    docx_bytes = _make_docx(content)
    chunks = parse_and_chunk(docx_bytes, "contract.docx")
    assert content in "".join(chunks)


def test_parse_xlsx_returns_cell_values():
    xlsx_bytes = _make_xlsx([["Name", "Score"], ["Alice", "95"], ["Bob", "87"]])
    chunks = parse_and_chunk(xlsx_bytes, "grades.xlsx")
    combined = "".join(chunks)
    assert "Name" in combined
    assert "Alice" in combined
    assert "95" in combined


def test_parse_xlsx_multiple_rows_joined():
    xlsx_bytes = _make_xlsx([["a", "b"], ["c", "d"]])
    chunks = parse_and_chunk(xlsx_bytes, "data.xlsx")
    combined = "".join(chunks)
    assert "a" in combined and "b" in combined
    assert "c" in combined and "d" in combined


def test_unsupported_extension_raises():
    with pytest.raises(ValueError, match="Unsupported file type"):
        parse_and_chunk(b"data", "file.csv")


def test_extension_case_insensitive():
    data = "uppercase extension".encode()
    chunks = parse_and_chunk(data, "FILE.TXT")
    assert "uppercase extension" in "".join(chunks)


# ---------------------------------------------------------------------------
# Chunking: text is split into exact 1000-char chunks
# ---------------------------------------------------------------------------

def test_chunk_size_exact_multiple():
    text = "x" * 3000
    result = _chunk(text)
    assert len(result) == 3
    assert all(len(c) == CHUNK_SIZE for c in result)


def test_chunk_size_non_multiple():
    text = "y" * 2500
    result = _chunk(text)
    assert len(result) == 3
    assert len(result[0]) == CHUNK_SIZE
    assert len(result[1]) == CHUNK_SIZE
    assert len(result[2]) == 500


def test_chunk_preserves_all_content():
    text = "a" * 1234
    result = _chunk(text)
    assert "".join(result) == text


def test_chunk_single_chunk_when_short():
    text = "short"
    result = _chunk(text)
    assert len(result) == 1
    assert result[0] == "short"


def test_parse_and_chunk_produces_correct_chunk_sizes():
    long_text = "z" * 2500
    chunks = parse_and_chunk(long_text.encode(), "data.txt")
    assert len(chunks) == 3
    assert len(chunks[0]) == CHUNK_SIZE
    assert len(chunks[1]) == CHUNK_SIZE
    assert len(chunks[2]) == 500


def test_chunk_empty_string_returns_one_empty_chunk():
    result = _chunk("")
    assert result == [""]


# ---------------------------------------------------------------------------
# Sanitization: control characters are removed
# ---------------------------------------------------------------------------

def test_sanitize_removes_null_bytes():
    assert _sanitize("he\x00llo") == "hello"


def test_sanitize_removes_bell_and_backspace():
    assert _sanitize("a\x07b\x08c") == "abc"


def test_sanitize_removes_delete_char():
    assert _sanitize("clean\x7ftext") == "cleantext"


def test_sanitize_preserves_tab_newline_cr():
    text = "col1\tcol2\nrow2\r\n"
    assert _sanitize(text) == text


def test_sanitize_preserves_printable_ascii():
    text = "Hello, World! 123 #$%"
    assert _sanitize(text) == text


def test_parse_and_chunk_strips_control_chars_in_txt():
    data = "clean\x00\x01\x1ftext".encode()
    chunks = parse_and_chunk(data, "file.txt")
    combined = "".join(chunks)
    assert "\x00" not in combined
    assert "\x01" not in combined
    assert "\x1f" not in combined
    assert "cleantext" in combined


def test_parse_and_chunk_strips_control_chars_in_docx():
    docx_bytes = _make_docx("valid content")
    chunks = parse_and_chunk(docx_bytes, "doc.docx")
    combined = "".join(chunks)
    for char in combined:
        code = ord(char)
        assert not (0x00 <= code <= 0x08), f"Found control char \\x{code:02x}"
        assert code not in (0x0b, 0x0c), f"Found control char \\x{code:02x}"
        assert not (0x0e <= code <= 0x1f), f"Found control char \\x{code:02x}"
        assert code != 0x7f, "Found DEL char"
