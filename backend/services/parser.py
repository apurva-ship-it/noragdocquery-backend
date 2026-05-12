import io
import re
from pathlib import Path

import fitz  # PyMuPDF
import docx  # python-docx
import openpyxl

CHUNK_SIZE = 1000
SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".xlsx", ".txt"})

# Matches control characters except horizontal tab (\x09), newline (\x0a), carriage return (\x0d)
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def _sanitize(text: str) -> str:
    return _CONTROL_RE.sub("", text)


def _chunk(text: str) -> list[str]:
    return [text[i : i + CHUNK_SIZE] for i in range(0, max(len(text), 1), CHUNK_SIZE)]


def _parse_pdf(data: bytes) -> str:
    doc = fitz.open(stream=data, filetype="pdf")
    return "\n".join(page.get_text() for page in doc)


def _parse_docx(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _parse_xlsx(data: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            lines.append("\t".join("" if cell is None else str(cell) for cell in row))
    return "\n".join(lines)


def _parse_txt(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


_PARSERS = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
    ".txt": _parse_txt,
}


def parse_and_chunk(data: bytes, filename: str) -> list[str]:
    ext = Path(filename).suffix.lower()
    parser = _PARSERS.get(ext)
    if parser is None:
        raise ValueError(f"Unsupported file type: {ext!r}")
    raw = parser(data)
    clean = _sanitize(raw)
    return _chunk(clean)
